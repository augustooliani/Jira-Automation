"""
Orquestra o fluxo completo (modo UI - abre o Edge para voce acompanhar):
  1. Le a BOM (Excel) e filtra linhas com status STATUS_FILTER.
  2. Para cada linha, localiza o datasheet correspondente ao part number.
  3. Cria a issue no Jira pela tela.
  4. Anexa o datasheet (PDF) e recarrega a tela.
  5. Completa a issue preenchendo todos os customfields da BOM.
  6. Substitui o "need erf" / "ok need erf" da coluna config.COLUMN_STATUS
     pela key da issue (ex: MAT-1234) com hyperlink para o Jira.
O navegador permanece aberto ao final para voce conferir os tickets.
"""

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

import config
from step2_find_datasheet import find_datasheet
from step3_create_jira_ui import (
    create_minimal_issue_ui,
    attach_file_ui,
    update_issue_fields_ui,
    set_assignee_ui,
    transition_issue_ui,
    start_browser,
)
from sharepoint_io import download_bom, upload_bom


def load_filtered_bom() -> pd.DataFrame:
    df = pd.read_excel(config.BOM_FILE, sheet_name=config.BOM_SHEET, engine="openpyxl")
    df[config.COLUMN_STATUS] = (
        df[config.COLUMN_STATUS].fillna("").astype(str).str.strip().str.lower()
    )
    return df[df[config.COLUMN_STATUS].str.contains(config.STATUS_FILTER, na=False)].copy()


def process_row(row, datasheet_root: Path) -> str | None:
    """Processa uma linha. Retorna a issue_key criada (ou None em caso de erro)."""
    part_number = str(row[config.COLUMN_PART_NUMBER]).strip()
    line = row[config.COLUMN_LINE]

    print("-" * 70)
    print(f"Linha BOM: {line} | Part Number: {part_number}")

    datasheet = find_datasheet(part_number, datasheet_root)
    if datasheet:
        print(f"  Datasheet localizado: {datasheet.name}")
    else:
        print("  [AVISO] datasheet nao encontrado")

    try:
        issue_key = create_minimal_issue_ui(row)
        print(f"  Issue criada: {issue_key}")
    except Exception as exc:
        print(f"  [ERRO] criando issue: {exc}")
        return None

    if datasheet:
        try:
            attach_file_ui(issue_key, datasheet)
            print(f"  Datasheet anexado em {issue_key}")
        except Exception as exc:
            print(f"  [ERRO] anexando datasheet: {exc}")

    try:
        update_issue_fields_ui(issue_key, row)
        print(f"  Campos preenchidos em {issue_key}")
    except Exception as exc:
        print(f"  [ERRO] preenchendo campos de {issue_key}: {exc}")

    # IMPORTANTE: a transicao precisa ser feita ANTES do assignee, porque
    # o workflow do Jira reatribui automaticamente o ticket ao mudar para
    # "Purchasing". Se setassemos antes, o valor seria sobrescrito.
    if getattr(config, "TEST_MODE", False):
        print(f"  [TEST_MODE] pulando transicao e assignee em {issue_key}")
    else:
        try:
            transition_issue_ui(issue_key, config.POST_CREATE_TRANSITION)
            print(f"  Status transicionado em {issue_key}")
        except Exception as exc:
            print(f"  [ERRO] transicionando {issue_key}: {exc}")

        try:
            set_assignee_ui(issue_key, config.POST_CREATE_ASSIGNEE_DISPLAY)
            print(f"  Assignee definido em {issue_key}")
        except Exception as exc:
            print(f"  [ERRO] definindo assignee de {issue_key}: {exc}")

    return issue_key


def write_keys_to_bom(keys_by_index: dict[int, str]) -> None:
    """Substitui o conteudo da coluna de status (R&D Remarks) pelo MAT-xxxx
    correspondente, com hyperlink apontando para a issue no Jira.
    Preserva todas as outras abas e formatacoes.
    keys_by_index: {indice_da_linha_no_dataframe: issue_key}
    """
    if not keys_by_index:
        print("Nenhuma issue criada -> planilha nao foi alterada.")
        return

    wb = load_workbook(config.BOM_FILE)
    ws = wb[config.BOM_SHEET]

    # Lê o cabecalho (linha 1) -> dicionario nome -> coluna (1-indexada)
    header = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value is not None:
            header[str(cell.value).strip()] = col_idx

    if config.COLUMN_STATUS not in header:
        raise RuntimeError(
            f"Coluna de status '{config.COLUMN_STATUS}' nao encontrada na planilha."
        )
    target_col = header[config.COLUMN_STATUS]

    link_font = Font(color="0563C1", underline="single")
    base_url = config.JIRA_BASE_URL.rstrip("/")

    # pandas leu com header=0, entao a linha do dataframe N corresponde
    # a linha (N + 2) no Excel (1 para o header + 1 porque eh 1-indexado).
    for df_idx, key in keys_by_index.items():
        excel_row = df_idx + 2
        cell = ws.cell(row=excel_row, column=target_col)
        cell.value = key
        cell.hyperlink = f"{base_url}/browse/{key}"
        cell.font = link_font

    wb.save(config.BOM_FILE)
    print(
        f"Planilha atualizada: {len(keys_by_index)} key(s) gravada(s) na coluna "
        f"'{config.COLUMN_STATUS}' de '{config.BOM_FILE}'."
    )


def main() -> None:
    # 0) Pergunta o link do SharePoint, baixa a BOM e aponta config.BOM_FILE para ela.
    print("Cole o link do SharePoint para a BOM (Enter sem nada para usar o arquivo local "
          f"'{config.BOM_FILE}'):")
    share_url = input("URL: ").strip()
    if share_url:
        start_browser()  # garante perfil/login antes de tentar baixar
        local_bom = Path("input") / "BOM_from_sharepoint.xlsx"
        download_bom(share_url, local_bom)
        config.BOM_FILE = str(local_bom)
    else:
        print(f"Usando arquivo local: {config.BOM_FILE}")

    datasheet_root = Path(config.DATASHEET_ROOT)
    filtered = load_filtered_bom()
    print(f"Total de linhas a processar: {len(filtered)}")

    start_browser()  # garante login antes do loop

    keys_by_index: dict[int, str] = {}
    seen_mpn: dict[str, str] = {}  # mpn -> issue_key ja criada
    for df_idx, row in filtered.iterrows():
        mpn = str(row[config.COLUMN_PART_NUMBER]).strip()
        mpn_key = mpn.lower() if mpn and mpn.lower() != "nan" else ""
        if mpn_key and mpn_key in seen_mpn:
            key = seen_mpn[mpn_key]
            print(f"  MPN duplicado ({mpn}) -> reutilizando ticket {key}")
            keys_by_index[df_idx] = key
        else:
            key = process_row(row, datasheet_root)
            if key:
                keys_by_index[df_idx] = key
                if mpn_key:
                    seen_mpn[mpn_key] = key

    # Resumo no console
    print("\n" + "=" * 70)
    print("RESUMO DAS ISSUES")
    print("=" * 70)
    if keys_by_index:
        first_occurrence: dict[str, int] = {}
        for df_idx, key in keys_by_index.items():
            if key not in first_occurrence:
                first_occurrence[key] = df_idx
        for df_idx, key in keys_by_index.items():
            line = filtered.loc[df_idx, config.COLUMN_LINE]
            pn = filtered.loc[df_idx, config.COLUMN_PART_NUMBER]
            tag = " [DUPLICADO]" if first_occurrence[key] != df_idx else ""
            print(f"  {key:<12}  Linha {line:<8}  {pn}{tag}")
    else:
        print("  (nenhuma issue criada)")

    # Atualiza a planilha local com as keys
    try:
        write_keys_to_bom(keys_by_index)
    except Exception as exc:
        print(f"[ERRO] gravando keys na planilha: {exc}")

    # Sobe a planilha de volta para o SharePoint (se veio de la)
    if share_url and keys_by_index:
        try:
            upload_bom(share_url, Path(config.BOM_FILE))
        except Exception as exc:
            print(f"[ERRO] subindo a BOM para o SharePoint: {exc}")
    elif share_url:
        print("Nenhuma key gerada -> upload para o SharePoint nao realizado.")

    # Navegador permanece aberto - encerra so quando o usuario quiser
    print("\nO navegador continua aberto para conferencia.")
    print("Pressione ENTER no terminal para encerrar...")
    try:
        input()
    except EOFError:
        pass


if __name__ == "__main__":
    main()
